"""ウィローズグループ CF現在地・6ヶ月予測の集計スクリプト。

使い方:
  python3 cf_forecast.py --wl wl.xlsx --bu bu.xlsx --balances balances.json [--months 6]

balances.json は「決算書・財務諸表等」の普通預金 補助残高一覧表から拾った月末残高(円)。
  {"base_month": "2026-09", "WL": 185243562, "BU": 280688976, "CR": 106522562,
   "history": {"2026-05": {"WL": ..., "BU": ..., "CR": ...}, ...}}

予測ロジック:
  - WL: PL・CF予実_WL の最新「キャッシュフロー（YYYY.MM~）」シートの「月末時の残高」行
  - BU: BU_CF.xlsx の最新「資金繰り表」シートの「翌月繰越金」行(定期預金は含めない)
  - 計画の基準月残高と実績の差(上振れ/下振れ)を以降の月にそのまま乗せる(リベース)
  - CR: 予測シートが無いため基準月残高を横置き
"""
import argparse
import datetime as dt
import json
import re

import openpyxl


def month_key(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}"
    return None


def latest_sheet(wb, pattern):
    cands = [ws.title for ws in wb.worksheets if re.search(pattern, ws.title)]
    if not cands:
        raise SystemExit(f"sheet not found: {pattern}")
    return cands[-1]  # シートは時系列順に右へ追加されている前提


def find_row(rows, label):
    for r in rows:
        if any(isinstance(c, str) and label in c for c in r[:3]):
            return r
    raise SystemExit(f"row not found: {label}")


def wl_plan(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = latest_sheet(wb, r"^キャッシュフロー（\d{4}\.\d+~）")
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    header = next(r for r in rows if sum(isinstance(c, dt.datetime) for c in r) >= 6)
    bal, sk, yc = find_row(rows, "月末時の残高"), find_row(rows, "預り敷金"), find_row(rows, "預り家賃")
    plan = {}
    for i, h in enumerate(header):
        m = month_key(h)
        if m and isinstance(bal[i], (int, float)):
            plan[m] = {"cash": bal[i], "deposits": (sk[i] or 0) + (yc[i] or 0)}
    return name, plan


def bu_plan(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = latest_sheet(wb, r"資金繰り表")
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    header = next(r for r in rows if sum(isinstance(c, dt.datetime) for c in r) >= 6)
    end = find_row(rows, "翌月繰越金")
    plan = {}
    for i, h in enumerate(header):
        m = month_key(h)
        if m and isinstance(end[i], (int, float)):
            plan[m] = {"cash": end[i]}  # 同じ月が複数列ある場合は右側(最新)を採用
    return name, plan


def next_months(base, n):
    y, m = map(int, base.split("-"))
    out = []
    for _ in range(n):
        m += 1
        if m > 12:
            y, m = y + 1, 1
        out.append(f"{y:04d}-{m:02d}")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--wl", required=True)
    ap.add_argument("--bu", required=True)
    ap.add_argument("--balances", required=True)
    ap.add_argument("--months", type=int, default=6)
    a = ap.parse_args()

    bal = json.load(open(a.balances, encoding="utf-8"))
    base = bal["base_month"]
    wl_sheet, wl = wl_plan(a.wl)
    bu_sheet, bu = bu_plan(a.bu)
    months = next_months(base, a.months)

    wl_gap = bal["WL"] - wl[base]["cash"]
    bu_gap = bal["BU"] - bu[base]["cash"]
    rows = []
    for m in months:
        w = wl.get(m, {}).get("cash")
        b = bu.get(m, {}).get("cash")
        w = None if w is None else w + wl_gap
        b = None if b is None else b + bu_gap
        c = bal["CR"]
        total = None if None in (w, b) else w + b + c
        rows.append({"month": m, "WL": w, "WL_usable": None if w is None else w - wl[m]["deposits"],
                     "BU": b, "CR": c, "total": total})

    out = {
        "base_month": base,
        "sources": {"WL": wl_sheet, "BU": bu_sheet},
        "actual": {k: bal[k] for k in ("WL", "BU", "CR")} | {"total": bal["WL"] + bal["BU"] + bal["CR"]},
        "plan_at_base": {"WL": wl[base]["cash"], "BU": bu[base]["cash"]},
        "gap_vs_plan": {"WL": wl_gap, "BU": bu_gap},
        "WL_usable_at_base": bal["WL"] - wl[base]["deposits"],
        "forecast": rows,
        "history": bal.get("history", {}),
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
