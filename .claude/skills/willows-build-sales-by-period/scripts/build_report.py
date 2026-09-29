#!/usr/bin/env python3
"""ウィローズビルド BU_CF「入力シート」→ 期別・種別別 販売分析ブック

usage: python build_report.py <BU_CF.xlsx> <output.xlsx> [基準日 YYYY-MM-DD]
"""
import sys
from collections import Counter
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

SRC_SHEET = "入力シート"
FIRST_ROW = 5
# 1-indexed 列番号（入力シート）
C_TYPE, C_NAME, C_BUY_SETTLE, C_SELL_CONTRACT, C_SELL_SETTLE, C_PRICE = 2, 3, 5, 6, 7, 8
C_BUY_COST = [9, 10, 11, 12, 13]                         # 手付金/契約時諸費用/残金/決済時諸費用/その他
C_PROD_COST = [16, 18, 20, 22, 24, 26, 28, 30, 32, 34, 36]  # リフォーム/解体/立退/建物手付・中間・残金/その他①〜⑤

# ビルド期（=グループ期−2）。12期は決算期変更の3ヶ月変則決算
PERIODS = [
    ("9期", date(2021, 7, 1), date(2022, 6, 30), ""),
    ("10期", date(2022, 7, 1), date(2023, 6, 30), ""),
    ("11期", date(2023, 7, 1), date(2024, 6, 30), ""),
    ("12期", date(2024, 7, 1), date(2024, 9, 30), "決算期変更の3ヶ月変則"),
    ("13期", date(2024, 10, 1), date(2025, 9, 30), ""),
    ("14期", date(2025, 10, 1), date(2026, 9, 30), ""),
    ("15期", date(2026, 10, 1), date(2027, 9, 30), ""),
    ("16期", date(2027, 10, 1), date(2028, 9, 30), ""),
]
BAND = 50_000_000  # 価格分布の刻み（5,000万円）

FONT = "Meiryo"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=14, bold=True)
FIN = Font(name=FONT, size=10, color="0000FF")
FNOTE = Font(name=FONT, size=9, color="666666")
FILL_H = PatternFill("solid", fgColor="1F4E78")
FILL_SUB = PatternFill("solid", fgColor="DDEBF7")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '#,##0;▲#,##0;"-"'
CNT = '#,##0;▲#,##0;"-"'
PCT = '0.0%;▲0.0%;"-"'
DT = "yyyy/mm/dd"


def yen_label(expr):
    """円の式 → 「5,000万」「1.5億」表記の式"""
    return f'IF({expr}>=100000000,TEXT({expr}/100000000,"0.0")&"億",TEXT({expr}/10000,"#,##0")&"万")'


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 40000 <= v <= 60000:
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True)[SRC_SHEET]
    rows = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        name = ws.cell(r, C_NAME).value
        if name in (None, ""):
            continue
        g = lambda c: ws.cell(r, c).value
        notes = []
        ptype = (str(g(C_TYPE)).strip() if g(C_TYPE) else "") or "未分類"
        if ptype == "未分類":
            notes.append("物件種別が空欄")
        buy = to_date(g(C_BUY_SETTLE))
        sell_c = to_date(g(C_SELL_CONTRACT))
        sell = to_date(g(C_SELL_SETTLE))
        for c, label in ((C_BUY_SETTLE, "仕入決済日"), (C_SELL_SETTLE, "売却決済日")):
            if g(c) not in (None, "") and to_date(g(c)) is None and not str(g(c)).startswith("00:00"):
                notes.append(f"{label}が日付でない({g(c)})")
        price = num(g(C_PRICE)) or None
        buy_cost = sum(num(g(c)) for c in C_BUY_COST)
        prod_cost = sum(num(g(c)) for c in C_PROD_COST)
        if sell and not price:
            notes.append("売却決済日ありだが販売価格が空欄")
        if sell and not buy:
            notes.append("仕入決済日が空欄→保有日数算出不可")
        if not buy_cost:
            notes.append("仕入原価未入力→粗利対象外")
        elif price and price < buy_cost + prod_cost:
            notes.append("原価が売価を上回る")
        rows.append(dict(no=g(1), type=ptype, name=str(name).strip(), buy=buy, sell_c=sell_c, sell=sell,
                         price=price, buy_cost=buy_cost or None, prod_cost=prod_cost, src_row=r,
                         note=" / ".join(notes)))
    return rows


def hdr(ws, row, col, labels, widths=None):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.border = FH, FILL_H, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[L(col + i)].width = w


def cell(ws, r, c, v, fmt=None, font=F, fill=None):
    x = ws.cell(r, c, v)
    x.font, x.border = font, BOX
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x


def build(src, out, base):
    rows = load(src)
    sold = sorted([r for r in rows if r["sell"]], key=lambda r: r["sell"])
    unsold = [r for r in rows if not r["sell"]]
    rows = sold + unsold
    types = [t for t, _ in Counter(r["type"] for r in sold).most_common()]
    first_sale = min(r["sell"] for r in sold)
    last_sale = max(r["sell"] for r in sold)
    periods = [p for p in PERIODS if p[2] >= first_sale and p[1] <= last_sale]
    max_price = max((r["price"] or 0) for r in sold)
    nbands = int(max_price // BAND) + 1

    wb = openpyxl.Workbook()

    # ---------------- 設定 ----------------
    st = wb.active
    st.title = "設定"
    st["A1"] = "設定（黄色セルのみ編集）"; st["A1"].font = FT
    cell(st, 3, 1, "基準日", font=FB); cell(st, 3, 2, base, DT, FIN, FILL_IN)
    cell(st, 3, 3, "これ以前の売却決済＝実績、以降＝予定", font=FNOTE)
    cell(st, 4, 1, "集計対象", font=FB); cell(st, 4, 2, "全て", None, FIN, FILL_IN)
    cell(st, 4, 3, "「全て」＝実績＋予定 ／「実績のみ」＝基準日以前の決済分だけ", font=FNOTE)
    dv = DataValidation(type="list", formula1='"全て,実績のみ"', allow_blank=False)
    st.add_data_validation(dv); dv.add("B4")
    cell(st, 5, 1, "価格帯の刻み(円)", font=FB); cell(st, 5, 2, BAND, YEN, FIN, FILL_IN)
    hdr(st, 7, 1, ["期", "開始日", "終了日", "備考"], [18, 14, 14, 60])
    P0 = 8
    for i, (name, s, e, note) in enumerate(PERIODS):
        r = P0 + i
        cell(st, r, 1, name, font=FIN, fill=FILL_IN)
        cell(st, r, 2, s, DT, FIN, FILL_IN); cell(st, r, 3, e, DT, FIN, FILL_IN)
        cell(st, r, 4, note or None)
    PN = P0 + len(PERIODS) - 1
    n = PN + 2
    for t in [
        "【前提】",
        "・期は売却決済日（売上計上日）で判定。ビルド期数＝グループ期数−2（ビルド15期＝2026.10–2027.9）",
        "・売上高＝販売価格（入力シートH列）。粗利＝販売価格−仕入原価−商品化原価（販売諸費用・支払利息は含めない＝会計上の粗利）",
        "・仕入原価＝入力シートI〜M列、商品化原価＝P,R,T,…,AJ列（リフォーム・解体・立退・建物・その他①〜⑤）",
        "・仕入原価が未入力の物件は粗利・粗利率の集計から除外（売上高・棟数には含む）",
        "・保有日数＝売却決済日−仕入決済日",
        f"・出所：{src.split('/')[-1]}「{SRC_SHEET}」 作成日 {date.today():%Y/%m/%d}",
    ]:
        st.cell(n, 1, t).font = FNOTE if t != "【前提】" else FB; n += 1
    BASE, MODE, BANDC = "設定!$B$3", "設定!$B$4", "設定!$B$5"
    PNAME, PSTART = f"設定!$A${P0}:$A${PN}", f"設定!$B${P0}:$B${PN}"

    # ---------------- 物件データ ----------------
    d = wb.create_sheet("物件データ")
    d["A1"] = "物件データ（入力シートから抽出・売却決済日順）"; d["A1"].font = FT
    d["A2"] = "青字＝入力シートからの転記値、黒字＝数式。売却決済日が空欄の物件は末尾（未売却）"; d["A2"].font = FNOTE
    heads = ["No", "物件種別", "物件名", "仕入決済日", "売却契約日", "売却決済日", "販売価格(円)",
             "仕入原価(円)", "商品化原価(円)", "期", "状態", "集計対象", "粗利(円)", "粗利率", "保有日数",
             "価格帯", "入力シート行", "データ注記"]
    hdr(d, 4, 1, heads, [6, 18, 30, 12, 12, 12, 15, 15, 15, 8, 8, 8, 15, 9, 9, 16, 9, 44])
    D0 = 5
    for i, x in enumerate(rows):
        r = D0 + i
        for c, v, f in [(1, x["no"], None), (2, x["type"], None), (3, x["name"], None), (4, x["buy"], DT),
                        (5, x["sell_c"], DT), (6, x["sell"], DT), (7, x["price"], YEN), (8, x["buy_cost"], YEN),
                        (9, x["prod_cost"], YEN), (17, x["src_row"], None), (18, x["note"] or None, None)]:
            cell(d, r, c, v, f, FIN if c != 18 else FNOTE)
        cell(d, r, 10, f'=IF(F{r}="","",IFERROR(INDEX({PNAME},MATCH(F{r},{PSTART},1)),"対象外"))')
        cell(d, r, 11, f'=IF(F{r}="","未売却",IF(F{r}<={BASE},"実績","予定"))')
        cell(d, r, 12, f'=IF(F{r}="",0,IF(OR({MODE}="全て",K{r}="実績"),1,0))')
        cell(d, r, 13, f'=IF(OR(G{r}="",H{r}=""),"",G{r}-H{r}-I{r})', YEN)
        cell(d, r, 14, f'=IF(M{r}="","",M{r}/G{r})', PCT)
        cell(d, r, 15, f'=IF(OR(D{r}="",F{r}=""),"",F{r}-D{r})', CNT)
        cell(d, r, 16, f'=IF(G{r}="","",' + yen_label(f"INT(G{r}/{BANDC})*{BANDC}") + '&"〜")')
    DN = D0 + len(rows) - 1
    d.freeze_panes = "D5"
    d.auto_filter.ref = f"A4:R{DN}"
    col = lambda c: f"物件データ!${c}${D0}:${c}${DN}"
    TY, PER, FLG, PRICE, GP, HOLD, COST = (col(c) for c in "BJLGMOH")

    def period_cols(ws, row, first_label, extra=("合計",), w0=22):
        hdr(ws, row, 1, [first_label] + [p[0] for p in periods] + list(extra),
            [w0] + [15] * (len(periods) + len(extra)))

    # ---------------- ①期別×種別 棟数 ----------------
    s1 = wb.create_sheet("①期別×種別")
    s1["A1"] = "① 期ごとの物件種別（売却決済ベース 棟数）"; s1["A1"].font = FT
    s1["A2"] = '=" 集計対象："&設定!B4&"／基準日："&TEXT(設定!B3,"yyyy/mm/dd")'; s1["A2"].font = FNOTE
    period_cols(s1, 4, "物件種別")
    for i, t in enumerate(types):
        r = 5 + i
        cell(s1, r, 1, t, font=FB)
        for j in range(len(periods)):
            c = L(2 + j)
            cell(s1, r, 2 + j, f'=COUNTIFS({TY},$A{r},{PER},{c}$4,{FLG},1)', CNT)
        cell(s1, r, 2 + len(periods), f"=SUM(B{r}:{L(1 + len(periods))}{r})", CNT, FB, FILL_TOT)
    tr = 5 + len(types)
    cell(s1, tr, 1, "合計", font=FB, fill=FILL_TOT)
    for j in range(len(periods) + 1):
        c = L(2 + j)
        cell(s1, tr, 2 + j, f"=SUM({c}5:{c}{tr - 1})", CNT, FB, FILL_TOT)
    # 構成比
    r0 = tr + 2
    s1.cell(r0, 1, "種別構成比（棟数）").font = FB
    period_cols(s1, r0 + 1, "物件種別")
    for i, t in enumerate(types):
        r = r0 + 2 + i
        cell(s1, r, 1, t, font=FB)
        for j in range(len(periods) + 1):
            c = L(2 + j)
            cell(s1, r, 2 + j, f"=IF({c}${tr}=0,\"\",{c}{5 + i}/{c}${tr})", PCT)
    s1.freeze_panes = "B5"

    # ---------------- ②種別別 期別KPI ----------------
    s2 = wb.create_sheet("②種別別KPI")
    s2["A1"] = "② 種別ごと・期ごとの売上高／粗利／保有日数／棟数"; s2["A1"].font = FT
    s2["A2"] = '=" 集計対象："&設定!B4&"／粗利・粗利率は仕入原価入力済の物件のみ"'; s2["A2"].font = FNOTE
    r = 4
    anchors = {}

    def block(title, fn, fmt, total_fn):
        nonlocal r
        s2.cell(r, 1, title).font = FB
        s2.cell(r, 1).fill = FILL_SUB
        period_cols(s2, r + 1, "物件種別")
        top = r + 2
        for i, t in enumerate(types + ["合計"]):
            rr = top + i
            is_tot = t == "合計"
            cell(s2, rr, 1, t, font=FB, fill=FILL_TOT if is_tot else None)
            for j in range(len(periods) + 1):
                pc = L(2 + j)
                pcrit = f"{pc}${r + 1}" if j < len(periods) else None
                tcrit = None if is_tot else f"$A{rr}"
                v = fn(tcrit, pcrit, rr, pc, top) if not is_tot else total_fn(pcrit, rr, pc, top)
                cell(s2, rr, 2 + j, v, fmt, FB if (is_tot or j == len(periods)) else F,
                     FILL_TOT if (is_tot or j == len(periods)) else None)
        anchors[title] = (top, top + len(types))
        r = top + len(types) + 2

    def crit(t, p, extra=""):
        parts = [f"{FLG},1"]
        if t:
            parts.append(f"{TY},{t}")
        if p:
            parts.append(f"{PER},{p}")
        else:
            parts.append(f'{PER},"?*"')
        return ",".join(parts) + extra

    count_fn = lambda t, p, rr, pc, top: f"=COUNTIFS({crit(t, p)})"
    block("棟数", count_fn, CNT, lambda p, rr, pc, top: f"=SUM({pc}{top}:{pc}{rr - 1})")
    sales_fn = lambda t, p, rr, pc, top: f"=SUMIFS({PRICE},{crit(t, p)})"
    block("売上高（円）", sales_fn, YEN, lambda p, rr, pc, top: f"=SUM({pc}{top}:{pc}{rr - 1})")
    gp_fn = lambda t, p, rr, pc, top: f"=SUMIFS({GP},{crit(t, p)})"
    block("粗利（円）※原価入力済のみ", gp_fn, YEN, lambda p, rr, pc, top: f"=SUM({pc}{top}:{pc}{rr - 1})")
    # 粗利率 = 粗利 / 原価入力済物件の売上
    def gpr_fn(t, p, rr, pc, top):
        c = crit(t, p, f',{COST},">0"')
        return f'=IFERROR(SUMIFS({GP},{c})/SUMIFS({PRICE},{c}),"")'
    block("粗利率 ※原価入力済のみ", gpr_fn, PCT, lambda p, rr, pc, top: gpr_fn(None, p, rr, pc, top))
    def hold_fn(t, p, rr, pc, top):
        c = crit(t, p, f',{HOLD},">=0"')
        return f'=IFERROR(AVERAGEIFS({HOLD},{c}),"")'
    block("平均保有日数（日）", hold_fn, '#,##0;▲#,##0;"-"', lambda p, rr, pc, top: hold_fn(None, p, rr, pc, top))
    def miss_fn(t, p, rr, pc, top):
        has_cost = crit(t, p, f',{COST},">0"')
        return f"=COUNTIFS({crit(t, p)})-COUNTIFS({has_cost})"
    block("参考：原価未入力で粗利対象外の棟数", miss_fn, CNT, lambda p, rr, pc, top: f"=SUM({pc}{top}:{pc}{rr - 1})")
    s2.freeze_panes = "B4"

    # ---------------- ③価格分布 ----------------
    s3 = wb.create_sheet("③成約価格分布")
    s3["A1"] = "③ 期ごとの平均成約価格と価格分布（5,000万円きざみ）"; s3["A1"].font = FT
    s3["A2"] = '=" 集計対象："&設定!B4&"／販売価格が空欄の物件は除外"'; s3["A2"].font = FNOTE
    period_cols(s3, 4, "指標", extra=("全期",))
    pc_crit = lambda pc: f"{PER},{pc}$4" if pc else f'{PER},"?*"'
    stats = [("棟数（価格あり）", lambda p: f'=COUNTIFS({FLG},1,{pc_crit(p)},{PRICE},">0")', CNT),
             ("平均成約価格（円）", lambda p: f'=IFERROR(AVERAGEIFS({PRICE},{FLG},1,{pc_crit(p)},{PRICE},">0"),"")', YEN),
             ("最低価格（円）", lambda p: f'=IFERROR(1/(1/_xlfn.MINIFS({PRICE},{FLG},1,{pc_crit(p)},{PRICE},">0")),"")', YEN),
             ("最高価格（円）", lambda p: f'=IFERROR(1/(1/_xlfn.MAXIFS({PRICE},{FLG},1,{pc_crit(p)},{PRICE},">0")),"")', YEN),
             ("売上高合計（円）", lambda p: f'=SUMIFS({PRICE},{FLG},1,{pc_crit(p)})', YEN)]
    for i, (lab, fn, fmt) in enumerate(stats):
        rr = 5 + i
        cell(s3, rr, 1, lab, font=FB)
        for j in range(len(periods) + 1):
            p = L(2 + j) if j < len(periods) else None
            cell(s3, rr, 2 + j, fn(p), fmt, FB if lab.startswith("平均") else F,
                 FILL_SUB if lab.startswith("平均") else None)
    rb = 5 + len(stats) + 2
    s3.cell(rb, 1, "価格帯別 棟数").font = FB
    hdr(s3, rb + 1, 1, ["価格帯"] + [p[0] for p in periods] + ["全期", "下限(円)", "上限(円)"])
    s3.column_dimensions[L(len(periods) + 3)].width = 14
    s3.column_dimensions[L(len(periods) + 4)].width = 14
    lo_c, hi_c = L(len(periods) + 3), L(len(periods) + 4)
    b0 = rb + 2
    for k in range(nbands):
        rr = b0 + k
        cell(s3, rr, len(periods) + 3, f"={BANDC}*{k}", YEN, FNOTE)
        cell(s3, rr, len(periods) + 4, f"={BANDC}*{k + 1}", YEN, FNOTE)
        cell(s3, rr, 1, (f'="〜"&{yen_label(f"{hi_c}{rr}")}&"未満"' if k == 0 else f'={yen_label(f"{lo_c}{rr}")}&"〜"&{yen_label(f"{hi_c}{rr}")}&"未満"'), font=FB)
        for j in range(len(periods) + 1):
            p = L(2 + j) if j < len(periods) else None
            lo = f'">="&${lo_c}{rr}' if k else '">0"'
            cell(s3, rr, 2 + j, f'=COUNTIFS({FLG},1,{pc_crit(p)},{PRICE},{lo},{PRICE},"<"&${hi_c}{rr})', CNT)
    bn = b0 + nbands - 1
    cell(s3, bn + 1, 1, "合計", font=FB, fill=FILL_TOT)
    for j in range(len(periods) + 1):
        c = L(2 + j)
        cell(s3, bn + 1, 2 + j, f"=SUM({c}{b0}:{c}{bn})", CNT, FB, FILL_TOT)
    rc = bn + 3
    s3.cell(rc, 1, "価格帯別 構成比").font = FB
    hdr(s3, rc + 1, 1, ["価格帯"] + [p[0] for p in periods] + ["全期"])
    for k in range(nbands):
        rr = rc + 2 + k
        cell(s3, rr, 1, f"=A{b0 + k}", font=FB)
        for j in range(len(periods) + 1):
            c = L(2 + j)
            cell(s3, rr, 2 + j, f'=IF({c}${bn + 1}=0,"",{c}{b0 + k}/{c}${bn + 1})', PCT)
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "clustered", 0
    ch.title, ch.y_axis.title = "期別 成約価格分布（棟数）", "棟数"
    ch.add_data(Reference(s3, min_col=2, max_col=1 + len(periods), min_row=rb + 1, max_row=bn), titles_from_data=True)
    ch.set_categories(Reference(s3, min_col=1, min_row=b0, max_row=bn))
    ch.height, ch.width = 9, 22
    s3.add_chart(ch, f"A{rc + nbands + 4}")
    s3.freeze_panes = "B5"

    wb.move_sheet("設定", offset=4)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(out)
    print(f"sold={len(sold)} unsold={len(unsold)} types={types} periods={[p[0] for p in periods]} bands={nbands}")


if __name__ == "__main__":
    base = date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else date.today()
    build(sys.argv[1], sys.argv[2], base)
