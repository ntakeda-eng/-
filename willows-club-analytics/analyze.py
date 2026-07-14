#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ウィローズクラブ 紹介・リピート 週次分析エンジン
=================================================

「⭐紹介・リピートリスト」スプレッドシートの「16期入力用」タブを解析し、
以下を集計して Slack 投稿用のレポートテキストを生成する。

  1. 会員内の行動の実数        … 紹介 / リピート / 成約 などの件数
  2. 誰がどの位獲得したか       … 社員別 社内紹介ポイント ランキング
  3. 今期(16期)の着地見込み     … 月次ペースからの年間着地予測

前週スナップショット(JSON)を渡すと「前週比」を付けて出力する。
実行後に当週スナップショットを保存し、翌週の比較に使う。

使い方:
    python analyze.py <xlsx> [--prev snapshots/前週.json] [--save snapshots/当週.json]
                            [--today YYYY-MM-DD] [--json]

xlsx は Google Drive からエクスポートした
application/vnd.openxmlformats-officedocument.spreadsheetml.sheet 形式のファイル。
"""
import argparse
import datetime
import json
import sys
from collections import Counter, defaultdict

import openpyxl

# ─────────────────────────────────────────────────────────────
# 会計年度・シート設定 (16期 = 2025/10/1 - 2026/9/30)
#   ※ 期が変わったら FY_START / FY_END / SHEET を更新する
# ─────────────────────────────────────────────────────────────
FY_LABEL = "16期"
FY_START = datetime.date(2025, 10, 1)
FY_END = datetime.date(2026, 9, 30)
SHEET = "16期入力用"

# 「16期入力用」タブの列インデックス (0-based)
NO, ACQ, SEIYAKU, SHUBETSU, GENSEN, KAITO_ID, TAIO_SHUBETSU, JIGYO, \
    POINT_GETTER, TANTO, KYAKU, KIKKAKE, SHOKAI_P, F1, SEIYAKU_P, F2, \
    SHAIN_P, COUNT, BIKO, MIRAIE = range(20)


def to_date(x):
    return x.date() if isinstance(x, datetime.datetime) else None


def num(x):
    return float(x) if isinstance(x, (int, float)) else None


def load_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"エラー: タブ '{SHEET}' が見つかりません。存在するタブ: {wb.sheetnames}")
    ws = wb[SHEET]
    rows = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, (int, float)) and v >= 2600000:   # 16期の紹介NO
            rows.append([ws.cell(r, c).value for c in range(1, 21)])
    return rows


def analyze(rows, today):
    total = len(rows)

    # ── 1. 行動の実数 ──────────────────────────────
    by_shubetsu = Counter(str(r[SHUBETSU]).strip() for r in rows if r[SHUBETSU])
    by_taio = Counter(str(r[TAIO_SHUBETSU]).strip() for r in rows if r[TAIO_SHUBETSU])
    by_jigyo = Counter(str(r[JIGYO]).strip() for r in rows if r[JIGYO])
    seiyaku_done = sum(1 for r in rows if to_date(r[SEIYAKU]))
    member_actions = sum(1 for r in rows if num(r[KAITO_ID]))

    # クラブ会員(お客様)別 行動数・成約数
    member = defaultdict(lambda: {"name": "", "actions": 0, "seiyaku": 0})
    for r in rows:
        mid = num(r[KAITO_ID])
        if not mid:
            continue
        m = member[int(mid)]
        m["actions"] += 1
        if not m["name"] and r[GENSEN]:
            m["name"] = str(r[GENSEN]).strip()
        if to_date(r[SEIYAKU]):
            m["seiyaku"] += 1

    # ── 2. 誰がどの位獲得したか (社員別) ─────────────
    emp = defaultdict(lambda: {"points": 0.0, "shokai": 0, "repeat": 0,
                               "shanai_hoka": 0, "records": 0})
    for r in rows:
        getter = str(r[POINT_GETTER]).strip() if r[POINT_GETTER] else ""
        if not getter:
            continue
        e = emp[getter]
        e["records"] += 1
        p = num(r[SHAIN_P])
        if p:
            e["points"] += p
        sh = str(r[SHUBETSU]).strip() if r[SHUBETSU] else ""
        if sh == "紹介":
            e["shokai"] += 1
        elif sh == "リピート":
            e["repeat"] += 1
        if str(r[TAIO_SHUBETSU]).strip() == "社内紹介(他部門)":
            e["shanai_hoka"] += 1

    # ── ポイント発行量 ───────────────────────────
    club_shokai = sum(num(r[SHOKAI_P]) or 0 for r in rows)
    club_seiyaku = sum(num(r[SEIYAKU_P]) or 0 for r in rows)
    shain_total = round(sum(num(r[SHAIN_P]) or 0 for r in rows), 1)

    # ── 3. 月次トレンドと着地見込み (紹介取得日ベース) ──
    monthly = Counter()
    in_fy = 0
    for r in rows:
        dte = to_date(r[ACQ])
        if dte and FY_START <= dte <= FY_END:
            monthly[(dte.year, dte.month)] += 1
            in_fy += 1

    # 経過月数(当月は日割り)
    elapsed = (today.year - FY_START.year) * 12 + (today.month - FY_START.month) \
        + today.day / 30.0
    elapsed = max(elapsed, 0.1)
    run_rate = in_fy / elapsed
    forecast = round(run_rate * 12)

    # 直近3完了月ペース参考値
    recent_keys = sorted(
        [k for k in monthly if datetime.date(k[0], k[1], 1) <
         datetime.date(today.year, today.month, 1)])[-3:]
    recent = sum(monthly[k] for k in recent_keys)
    recent_rate = recent / len(recent_keys) if recent_keys else run_rate
    forecast_recent = round(in_fy + recent_rate * (12 - elapsed))

    return {
        "total_records": total,
        "in_fy_records": in_fy,
        "seiyaku_done": seiyaku_done,
        "seiyaku_rate": round(seiyaku_done / total * 100, 1) if total else 0,
        "member_actions": member_actions,
        "by_shubetsu": dict(by_shubetsu.most_common()),
        "by_taio": dict(by_taio.most_common()),
        "by_jigyo": dict(by_jigyo.most_common()),
        "club_shokai_p": club_shokai,
        "club_seiyaku_p": club_seiyaku,
        "shain_total_p": shain_total,
        "elapsed_months": round(elapsed, 2),
        "run_rate_per_month": round(run_rate, 1),
        "forecast_annual": forecast,
        "forecast_recent3": forecast_recent,
        "monthly": {f"{y}/{m:02d}": monthly[(y, m)] for (y, m) in sorted(monthly)},
        "emp_ranking": sorted(
            ({"name": k, **v} for k, v in emp.items()),
            key=lambda x: (-x["points"], -x["records"]))[:10],
        "member_ranking": sorted(
            ({"id": k, **v} for k, v in member.items()),
            key=lambda x: (-x["actions"], -x["seiyaku"]))[:10],
    }


def delta(cur, prev, key):
    """前週比の文字列 (例: '+3')。prev が無ければ空。"""
    if not prev or key not in prev:
        return ""
    d = round(cur - prev[key], 1)
    if d == 0:
        return " (±0)"
    d = int(d) if float(d).is_integer() else d
    return f" ({'+' if d > 0 else ''}{d})"


def format_report(a, prev, today):
    L = []
    L.append(f"📊 **ウィローズクラブ {FY_LABEL} 週次レポート** ({today:%Y/%m/%d}時点)")
    fy_pct = round(a["elapsed_months"] / 12 * 100)
    L.append(f"期の進捗: 第{int(a['elapsed_months'])+1}ヶ月目 / 12ヶ月 (約{fy_pct}%経過)")
    L.append("")

    # 1. 実数
    L.append("**① 会員内の行動の実数**")
    L.append(f"・紹介/リピート 総登録数: **{a['total_records']}件**"
             f"{delta(a['total_records'], prev, 'total_records')}")
    L.append(f"・成約完了: **{a['seiyaku_done']}件** (成約率 {a['seiyaku_rate']}%)"
             f"{delta(a['seiyaku_done'], prev, 'seiyaku_done')}")
    L.append(f"・当期発生(紹介取得日が{FY_LABEL}内): {a['in_fy_records']}件"
             f"{delta(a['in_fy_records'], prev, 'in_fy_records')}")
    sb = a["by_shubetsu"]
    L.append("・種別内訳: " + " / ".join(f"{k} {v}" for k, v in sb.items()))
    jg = list(a["by_jigyo"].items())[:5]
    L.append("・事業内訳(上位5): " + " / ".join(f"{k} {v}" for k, v in jg))
    L.append("")

    # 2. 獲得ランキング
    L.append("**② 誰がどの位獲得したか (社員別 社内紹介pt)**")
    prev_emp = {e["name"]: e for e in prev.get("emp_ranking", [])} if prev else {}
    for i, e in enumerate(a["emp_ranking"], 1):
        d = ""
        if e["name"] in prev_emp:
            dd = round(e["points"] - prev_emp[e["name"]]["points"], 1)
            if dd:
                dd = int(dd) if float(dd).is_integer() else dd
                d = f" ({'+' if dd > 0 else ''}{dd})"
        L.append(f"{i}. {e['name']}  **{round(e['points'],1)}pt**{d}"
                 f"  (紹介{e['shokai']}/ﾘﾋﾟｰﾄ{e['repeat']}/他部門{e['shanai_hoka']})")
    L.append(f"（社員pt総発行 {a['shain_total_p']}pt / "
             f"クラブpt発行 紹介相談{int(a['club_shokai_p']):,} + "
             f"成約{int(a['club_seiyaku_p']):,}）")
    L.append("")

    # 3. 着地見込み
    L.append(f"**③ {FY_LABEL} 着地見込み**")
    L.append(f"・月次ペース {a['run_rate_per_month']}件/月 → "
             f"**年間着地 約{a['forecast_annual']}件**")
    L.append(f"・直近3ヶ月ペース基準の着地: 約{a['forecast_recent3']}件")
    mk = list(a["monthly"].items())
    L.append("・月次推移: " + " ".join(f"{k[-2:]}月:{v}" for k, v in mk))
    L.append("")

    # 会員(お客様)側 上位
    L.append("**④ 紹介・リピートが多いクラブ会員(お客様) 上位5**")
    for i, m in enumerate(a["member_ranking"][:5], 1):
        nm = m["name"] or f"ID:{m['id']}"
        L.append(f"{i}. {nm}  行動{m['actions']}件 / 成約{m['seiyaku']}件")
    L.append("")
    L.append("_※成約は成約後入力のためタイムラグあり。直近月は今後上振れる見込み。_")
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--prev", help="前週スナップショットJSON")
    ap.add_argument("--save", help="当週スナップショットJSON保存先")
    ap.add_argument("--today", help="基準日 YYYY-MM-DD (省略時は本日)")
    ap.add_argument("--json", action="store_true", help="レポートでなくJSONを出力")
    args = ap.parse_args()

    today = (datetime.datetime.strptime(args.today, "%Y-%m-%d").date()
             if args.today else datetime.date.today())

    rows = load_rows(args.xlsx)
    a = analyze(rows, today)

    prev = None
    if args.prev:
        try:
            with open(args.prev, encoding="utf-8") as f:
                prev = json.load(f)
        except FileNotFoundError:
            pass

    if args.json:
        print(json.dumps(a, ensure_ascii=False, indent=2))
    else:
        print(format_report(a, prev, today))

    if args.save:
        a["_generated_at"] = today.isoformat()
        with open(args.save, "w", encoding="utf-8") as f:
            json.dump(a, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
